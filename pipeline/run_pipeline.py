import sys, os, json
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import anthropic
from dotenv import load_dotenv
load_dotenv()

client = anthropic.Anthropic(api_key=os.getenv("ANTHROPIC_API_KEY"))
MODEL = "claude-sonnet-4-20250514"

def analyse(transcript):
    r = client.messages.create(model=MODEL, max_tokens=2000,
        messages=[{"role":"user","content":"""Analyse this therapy session.
Return ONLY valid JSON with these fields:
presenting_themes, cognitive_distortions, therapeutic_alliance,
patient_engagement, disclosure_depth, resistance_indicators,
therapist_technique, rupture_repair, risk_flags,
session_quality_score, key_patient_quotes, summary.
TRANSCRIPT: """ + transcript["full_text"]}])
    clean = r.content[0].text.replace("```json","").replace("```","").strip()
    a = json.loads(clean)
    a["patient_id"] = transcript["patient_id"]
    a["date"] = transcript["date"]
    return a

def document(transcript, analysis, session_number):
    r = client.messages.create(model=MODEL, max_tokens=2000,
        messages=[{"role":"user","content":f"""
Generate a structured clinical note for EMR copy-paste.
DATE: {transcript['date']} | PATIENT: {transcript['patient_id']} | SESSION: {session_number}
Sections: PRESENTING COMPLAINT, MENTAL STATE EXAMINATION,
RISK ASSESSMENT, SESSION CONTENT, MANAGEMENT PLAN, NEXT SESSION FOCUS.
End with: *** AI-GENERATED — REQUIRES CLINICIAN REVIEW BEFORE FILING ***
TRANSCRIPT: {transcript['full_text'][:2000]}
ANALYSIS: {json.dumps(analysis)}"""}])
    note_text = r.content[0].text
    os.makedirs("outputs/clinical_notes", exist_ok=True)
    from docx import Document
    doc = Document()
    doc.add_heading("Clinical Session Note — CONFIDENTIAL", 0)
    for line in note_text.split("\n"):
        doc.add_paragraph(line)
    filepath = f"outputs/clinical_notes/{transcript['patient_id']}_{transcript['date']}_session{session_number}.docx"
    doc.save(filepath)
    print(f"Note saved: {filepath}")
    return {"note_text": note_text, "filepath": filepath}

def code(analysis):
    r = client.messages.create(model=MODEL, max_tokens=1000,
        messages=[{"role":"user","content":"""Suggest ICD-11 and DSM-5 codes.
Return ONLY valid JSON: icd11_primary, icd11_secondary, dsm5_primary,
dsm5_differential, coding_confidence, coding_notes,
requires_further_assessment, suggested_assessments.
ANALYSIS: """ + json.dumps(analysis)}])
    clean = r.content[0].text.replace("```json","").replace("```","").strip()
    return json.loads(clean)

def recommend(analysis):
    r = client.messages.create(model=MODEL, max_tokens=1000,
        messages=[{"role":"user","content":"""Provide supervision recommendations.
Return ONLY valid JSON: guided_questions_next_session, under_explored_areas,
therapy_technique_suggestions, therapy_modality_adjustment,
homework_suggestions, escalation_required, escalation_reason,
supervisor_review_recommended, overall_recommendation.
ANALYSIS: """ + json.dumps(analysis)}])
    clean = r.content[0].text.replace("```json","").replace("```","").strip()
    return json.loads(clean)

def medicolegal(note_text, analysis):
    r = client.messages.create(model=MODEL, max_tokens=1000,
        messages=[{"role":"user","content":"""Audit this clinical note.
Return ONLY valid JSON: consent_documented, risk_assessment_complete,
capacity_assessment_documented, safeguarding_addressed,
confidentiality_limits_documented, management_plan_present,
follow_up_documented, flags, overall_adequacy,
medicolegal_risk_level, recommended_actions.
NOTE: """ + note_text[:1000] + " ANALYSIS: " + json.dumps(analysis)}])
    clean = r.content[0].text.replace("```json","").replace("```","").strip()
    return json.loads(clean)

def run():
    transcript = {
        "patient_id":   "PT001",
        "therapist_id": "TH001",
        "date":         "2026-03-15",
        "full_text": """
        Therapist: How have you been feeling this week?
        Patient: Not great. Everything feels like my fault.
        I cannot sleep and feel hopeless most days.
        Therapist: Any thoughts of harming yourself?
        Patient: No, nothing like that. Just very low.
        Therapist: Let us explore those feelings further.
        """
    }

    print("\n── STAGE 1: Thematic Analysis ──────────")
    analysis = analyse(transcript)
    print(f"Quality score : {analysis.get('session_quality_score')}")
    print(f"Risk flags    : {analysis.get('risk_flags')}")

    print("\n── STAGE 2: Clinical Note ──────────────")
    note = document(transcript, analysis, 1)

    print("\n── STAGE 3: ICD-11 / DSM-5 Coding ─────")
    codes = code(analysis)
    print(f"ICD-11 : {codes.get('icd11_primary')}")
    print(f"DSM-5  : {codes.get('dsm5_primary')}")

    print("\n── STAGE 4: Recommendations ────────────")
    rec = recommend(analysis)
    print(f"Escalation required : {rec.get('escalation_required')}")

    print("\n── STAGE 5: Medicolegal Audit ──────────")
    audit = medicolegal(note["note_text"], analysis)
    print(f"Adequacy   : {audit.get('overall_adequacy')}")
    print(f"Risk level : {audit.get('medicolegal_risk_level')}")

    print("\n" + "="*50)
    print("HUMAN-IN-THE-LOOP CHECKPOINT")
    print("="*50)
    print(f"Patient    : {transcript['patient_id']}")
    print(f"Risk flags : {analysis.get('risk_flags')}")
    print(f"ICD-11     : {codes.get('icd11_primary')}")
    print(f"DSM-5      : {codes.get('dsm5_primary')}")
    print(f"Medicolegal: {audit.get('overall_adequacy')} — {audit.get('medicolegal_risk_level')} risk")
    print(f"Escalation : {rec.get('escalation_required')}")
    print(f"Note at    : {note['filepath']}")
    print("="*50)

    approved = input("\nApprove and file this record? (yes/no): ")
    if approved.strip().lower() != "yes":
        print("NOT filed. Returning for clinician revision.")
        return None
    print("\n✓ Pipeline complete. Record approved.")

if __name__ == "__main__":
    run()
