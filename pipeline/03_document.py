import anthropic, json, os
from docx import Document
from dotenv import load_dotenv
load_dotenv()

client = anthropic.Anthropic(api_key=os.getenv("ANTHROPIC_API_KEY"))
NOTES_DIR = "outputs/clinical_notes"

def generate_clinical_note(transcript, analysis, session_number=1):
    response = client.messages.create(
        model="claude-sonnet-4-20250514",
        max_tokens=2000,
        messages=[{"role": "user", "content": f"""
Generate a structured clinical note for EMR copy-paste.

DATE: {transcript['date']}
PATIENT ID: {transcript['patient_id']}
THERAPIST ID: {transcript['therapist_id']}
SESSION: {session_number}

Include these sections:
PRESENTING COMPLAINT:
MENTAL STATE EXAMINATION:
  Appearance and Behaviour:
  Speech:
  Mood (subjective):
  Affect (observed):
  Thought Form:
  Thought Content:
  Perception:
  Cognition:
  Insight:
  Judgement:
RISK ASSESSMENT:
  Self-harm:
  Harm to others:
  Safeguarding:
SESSION CONTENT:
RESPONSE TO THERAPY:
MANAGEMENT PLAN:
NEXT SESSION FOCUS:
CLINICIAN SIGNATURE: _________________ DATE: _________________
*** AI-GENERATED — REQUIRES CLINICIAN REVIEW BEFORE FILING ***

TRANSCRIPT: {transcript['full_text'][:2000]}
ANALYSIS: {json.dumps(analysis, indent=2)}
"""}]
    )
    note_text = response.content[0].text
    os.makedirs(NOTES_DIR, exist_ok=True)
    doc = Document()
    doc.add_heading("Clinical Session Note — CONFIDENTIAL", 0)
    for line in note_text.split("\n"):
        doc.add_paragraph(line)
    filename = f"{transcript['patient_id']}_{transcript['date']}_session{session_number}.docx"
    filepath = os.path.join(NOTES_DIR, filename)
    doc.save(filepath)
    print(f"Note saved: {filepath}")
    return {"note_text": note_text, "filepath": filepath}

if __name__ == "__main__":
    print("Documentation module ready.")
