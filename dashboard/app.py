import streamlit as st
import json, os, tempfile, anthropic, base64
import requests as _req
from dotenv import load_dotenv
from datetime import datetime

try:
    ANTHROPIC_API_KEY = st.secrets["ANTHROPIC_API_KEY"]
except Exception:
    load_dotenv()
    ANTHROPIC_API_KEY = os.getenv("ANTHROPIC_API_KEY")

if not ANTHROPIC_API_KEY:
    st.error("API key not found.")
    st.stop()

client = anthropic.Anthropic(api_key=ANTHROPIC_API_KEY)
MODEL = "claude-sonnet-4-20250514"

st.set_page_config(page_title="TherapyAI DEMO", page_icon="🧠", layout="centered")
st.markdown('<style>[data-testid="collapsedControl"]{display:none}section[data-testid="stSidebar"]{display:none}.block-container{padding:1rem;max-width:700px}</style>', unsafe_allow_html=True)

if "ai_count" not in st.session_state: st.session_state.ai_count = 0
if "results" not in st.session_state: st.session_state.results = {}

st.markdown("""<div style='background:linear-gradient(135deg,#1a1a2e,#0f3460);padding:22px 20px;border-radius:14px;margin-bottom:8px;'>
<span style='background:#e74c3c;color:white;font-size:0.7em;font-weight:700;padding:3px 10px;border-radius:20px;'>DEMO VERSION</span>
<h1 style='color:#fff;margin:8px 0 0 0;font-size:1.35em;font-weight:700;'>🧠 TherapyAI — Automated Clerking, Transcription & Risk Stratification</h1>
<p style='color:#a0c4ff;margin:6px 0 0 0;font-size:0.85em;'>by <strong>Dr Naim AI Team</strong> · Hospital Tengku Permaisuri Norashikin (HTPN)</p>
</div>""", unsafe_allow_html=True)

st.caption("⚙️ Demo version for feasibility and validity testing purposes only.")
c1, c2 = st.columns(2)
c1.metric("🤖 AI Analyses Run (this session)", st.session_state.ai_count)
c2.metric("📅 Today", datetime.now().strftime("%d %b %Y"))

st.warning("⚠️ **Disclaimer:** For automation of therapy clerking only. **No patient records stored.** Audio deleted after transcription. All outputs require clinician review. Does not replace clinical judgement.")

st.divider()
st.subheader("① Patient Details")
col1,col2 = st.columns(2)
with col1:
    patient_id   = st.text_input("Patient ID", value="PT001")
    therapist_id = st.text_input("Therapist ID", value="TH001")
with col2:
    session_num  = st.number_input("Session Number", min_value=1, value=1)
    session_date = st.date_input("Session Date", value=datetime.now())

st.divider()
st.subheader("② Session Input")
st.caption("Upload audio, image, or paste transcript — only one needed.")
t1,t2,t3 = st.tabs(["🎙️ Audio","🖼️ Image / Referral Letter","⌨️ Paste Transcript"])
transcript_text = ""; referral_context = ""

with t1:
    st.info("Supports .ogg .wav .mp3 .m4a .mp4")
    lang = st.radio("Language:", ["English","Bahasa Malaysia","Auto-detect"], horizontal=True)
    lmap = {"English":"en","Bahasa Malaysia":"ms","Auto-detect":None}
    af = st.file_uploader("Upload audio", type=["ogg","wav","mp3","m4a","mp4"], label_visibility="collapsed")
    if af:
        st.audio(af)
        if st.button("🎙️ Transcribe Now", use_container_width=True):
            with st.spinner("Transcribing..."):
                import whisper
                sx = "."+af.name.split(".")[-1]
                with tempfile.NamedTemporaryFile(delete=False,suffix=sx) as tmp:
                    tmp.write(af.read()); tp=tmp.name
                mw = whisper.load_model("base")
                res = mw.transcribe(tp,language=lmap[lang]) if lmap[lang] else mw.transcribe(tp)
                st.session_state["transcript"] = res["text"]
                os.unlink(tp)
            st.success(f"✅ Done · Language: {res.get('language','?')}")
    if "transcript" in st.session_state:
        st.session_state["transcript"] = st.text_area("Transcribed text (editable)", value=st.session_state["transcript"], height=180)
        transcript_text = st.session_state["transcript"]

with t2:
    st.info("Upload photo of referral letter, handwritten notes, or assessment form.")
    img = st.file_uploader("Upload image", type=["jpg","jpeg","png","webp"], label_visibility="collapsed")
    if img:
        st.image(img, use_column_width=True)
        if st.button("📄 Extract Clinical Info", use_container_width=True):
            with st.spinner("Extracting..."):
                ib = img.read(); ib64 = base64.standard_b64encode(ib).decode()
                mt = {"jpg":"image/jpeg","jpeg":"image/jpeg","png":"image/png","webp":"image/webp"}.get(img.name.split(".")[-1].lower(),"image/jpeg")
                rr = client.messages.create(model=MODEL,max_tokens=1500,messages=[{"role":"user","content":[
                    {"type":"image","source":{"type":"base64","media_type":mt,"data":ib64}},
                    {"type":"text","text":"Extract clinical info: DOCUMENT TYPE: PATIENT DETAILS: REFERRING CLINICIAN: REASON FOR REFERRAL: RELEVANT HISTORY: MEDICATIONS: INVESTIGATIONS: FINDINGS: RECOMMENDATIONS: Write Not stated if absent."}]}])
                st.session_state["referral_context"] = rr.content[0].text
            st.success("✅ Extracted")
    if "referral_context" in st.session_state:
        st.session_state["referral_context"] = st.text_area("Extracted info (editable)", value=st.session_state["referral_context"], height=200)
        referral_context = st.session_state["referral_context"]

with t3:
    transcript_text = st.text_area("Paste transcript here", height=200, placeholder="Therapist: How have you been?\nPatient: ...")

st.divider()
st.subheader("③ Run AI Analysis")
run_btn = st.button("▶ Run AI Analysis", type="primary", use_container_width=True)

def rc(prompt, max_tokens=2000):
    return client.messages.create(model=MODEL,max_tokens=max_tokens,messages=[{"role":"user","content":prompt}]).content[0].text

def pj(text):
    c = text.replace("```json","").replace("```","").strip()
    s=c.find("{"); e=c.rfind("}")+1
    if s==-1 or e==0: raise ValueError("No JSON")
    return json.loads(c[s:e])

ft = transcript_text or st.session_state.get("transcript","") or st.session_state.get("referral_context","") or referral_context

if run_btn and ft:
    st.session_state.ai_count += 1
    ds = session_date.strftime("%Y-%m-%d")
    ref = st.session_state.get("referral_context","") or referral_context
    tra = st.session_state.get("transcript","") or transcript_text
    ctx = ""
    if ref: ctx += f"\n\nREFERRAL:\n{ref}"
    if tra: ctx += f"\n\nTRANSCRIPT:\n{tra}"

    pg = st.progress(0, text="Starting...")
    pg.progress(10, text="Stage 1 — Thematic analysis...")
    analysis = pj(rc('Analyse therapy session. English or BM transcript, output English JSON only: {"presenting_themes":[],"cognitive_distortions":[],"therapeutic_alliance":0,"patient_engagement":0,"disclosure_depth":"","resistance_indicators":[],"therapist_technique":[],"rupture_repair":false,"risk_flags":[],"session_quality_score":0,"key_patient_quotes":[],"summary":""}\n'+ctx))

    pg.progress(30, text="Stage 2 — Clinical note...")
    note = rc(f"Structured clinical note for EMR. DATE:{ds} PATIENT:{patient_id} THERAPIST:{therapist_id} SESSION:{session_num}. Sections: PRESENTING COMPLAINT, MENTAL STATE EXAM (Appearance Speech Mood Affect Thought Form Thought Content Perception Cognition Insight Judgement), RISK ASSESSMENT (Self-harm Harm-to-others Safeguarding), SESSION CONTENT, RESPONSE TO THERAPY, MANAGEMENT PLAN, NEXT SESSION FOCUS. End: ***AI-GENERATED — CLINICIAN REVIEW REQUIRED***\n{ctx[:2500]}\nANALYSIS:{json.dumps(analysis)}")

    pg.progress(50, text="Stage 3 — ICD-11 / DSM-5...")
    codes = pj(rc('ICD-11 and DSM-5 codes. JSON only: {"icd11_primary":{"code":"","description":""},"icd11_secondary":[],"dsm5_primary":{"code":"","description":"","criteria_met":[]},"dsm5_differential":[],"coding_confidence":"","coding_notes":"","suggested_assessments":[]}\nANALYSIS:'+json.dumps(analysis), max_tokens=1000))

    pg.progress(70, text="Stage 4 — Recommendations...")
    rec = pj(rc('Supervision recommendations. JSON only: {"guided_questions_next_session":[],"therapy_technique_suggestions":[],"therapy_modality_adjustment":"","homework_suggestions":[],"escalation_required":false,"escalation_reason":null,"supervisor_review_recommended":false,"overall_recommendation":""}\nANALYSIS:'+json.dumps(analysis), max_tokens=1000))

    pg.progress(90, text="Stage 5 — Documentation quality...")
    audit = pj(rc('Documentation completeness audit. JSON only: {"consent_documented":false,"risk_assessment_complete":false,"safeguarding_addressed":false,"management_plan_present":false,"follow_up_documented":false,"missing_elements":[],"documentation_quality_score":0,"overall_completeness":"","standard_recommendations":[],"priority_actions":[]}\nNOTE:'+note[:800], max_tokens=1000))

    pg.progress(100, text="✅ Complete")
    st.session_state.results = {"analysis":analysis,"note":note,"codes":codes,"rec":rec,"audit":audit,"ds":ds}
    st.success("✅ Complete — review outputs below before filing")

elif run_btn:
    st.error("Please upload audio, upload image, or paste transcript first.")

if st.session_state.results:
    rv=st.session_state.results
    analysis=rv["analysis"]; note=rv["note"]; codes=rv["codes"]; rec=rv["rec"]; audit=rv["audit"]; ds=rv["ds"]

    st.divider(); st.subheader("④ Summary Metrics")
    m1,m2,m3=st.columns(3)
    m1.metric("Session Quality", f"{analysis.get('session_quality_score')}/10")
    m2.metric("Therapeutic Alliance", f"{analysis.get('therapeutic_alliance')}/5")
    m3.metric("Patient Engagement", f"{analysis.get('patient_engagement')}/5")
    m4,m5=st.columns(2)
    m4.metric("Doc Completeness", audit.get('overall_completeness','—').upper())
    m5.metric("Escalation","⚠️ YES" if rec.get('escalation_required') else "✅ NO")

    if analysis.get("risk_flags"): st.error(f"🚨 Risk Flags: {' · '.join(analysis['risk_flags'])}")
    else: st.success("✅ No risk flags")

    st.divider()
    tb1,tb2,tb3,tb4,tb5=st.tabs(["📋 Analysis","📄 Clinical Note","🏷️ ICD-11/DSM-5","💡 Recommendations","📝 Doc Quality"])

    with tb1:
        st.write("**Presenting Themes**");[st.write(f"• {t}") for t in analysis.get("presenting_themes",[])]
        st.write("**Cognitive Distortions**");[st.write(f"• {c}") for c in analysis.get("cognitive_distortions",[])]
        st.write(f"**Disclosure Depth:** {analysis.get('disclosure_depth','—')}")
        st.write("**Therapist Techniques**");[st.write(f"• {t}") for t in analysis.get("therapist_technique",[])]
        st.write("**Key Quotes**");[st.info(f'"{q}"') for q in analysis.get("key_patient_quotes",[])]
        st.write("**Summary**"); st.write(analysis.get("summary","—"))

    with tb2:
        st.warning("⚠️ Review and edit before filing in HIS/EMR")
        en = st.text_area("Clinical Note (editable)", value=note, height=400)
        d1,d2=st.columns(2)
        with d1:
            st.download_button("⬇️ Download .txt", data=en.encode(), file_name=f"{patient_id}_{ds}_s{session_num}.txt", mime="text/plain", use_container_width=True)
        with d2:
            from docx import Document as DD; from io import BytesIO
            dc=DD(); dc.add_heading("Clinical Note — CONFIDENTIAL",0); dc.add_paragraph(f"Patient:{patient_id} | Date:{ds} | Session:{session_num}")
            [dc.add_paragraph(l) for l in en.split("\n")]
            bf=BytesIO(); dc.save(bf); bf.seek(0)
            st.download_button("⬇️ Download .docx", data=bf.read(), file_name=f"{patient_id}_{ds}_s{session_num}.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document", use_container_width=True)
        st.info("💡 .txt → phone Notes app → copy-paste to HIS | .docx → Word/Google Docs")

    with tb3:
        st.warning("⚠️ AI-suggested — clinician must confirm")
        icd=codes.get("icd11_primary",{}); dsm=codes.get("dsm5_primary",{})
        st.write("**ICD-11**"); st.code(f"{icd.get('code','—')} {icd.get('description','—')}")
        st.write("**DSM-5**");  st.code(f"{dsm.get('code','—')} {dsm.get('description','—')}")
        st.write(f"Confidence: {codes.get('coding_confidence','—')} | Notes: {codes.get('coding_notes','—')}")
        if codes.get("suggested_assessments"): [st.write(f"• {a}") for a in codes["suggested_assessments"]]

    with tb4:
        st.write("**Guided Questions — Next Session**");[st.write(f"• {q}") for q in rec.get("guided_questions_next_session",[])]
        st.write("**Technique Suggestions**");[st.write(f"• {t}") for t in rec.get("therapy_technique_suggestions",[])]
        st.write(f"**Modality:** {rec.get('therapy_modality_adjustment','—')}")
        st.write("**Homework**");[st.write(f"• {h}") for h in rec.get("homework_suggestions",[])]
        st.info(rec.get("overall_recommendation","—"))
        if rec.get("escalation_required"): st.error(f"🚨 Escalate: {rec.get('escalation_reason','—')}")

    with tb5:
        st.info("Documentation Quality & Completeness Checker")
        st.metric("Quality Score", f"{audit.get('documentation_quality_score',0)}/10")
        c1,c2=st.columns(2)
        with c1:
            st.write(f"Consent: `{audit.get('consent_documented')}`")
            st.write(f"Risk assessment: `{audit.get('risk_assessment_complete')}`")
            st.write(f"Safeguarding: `{audit.get('safeguarding_addressed')}`")
        with c2:
            st.write(f"Management plan: `{audit.get('management_plan_present')}`")
            st.write(f"Follow-up: `{audit.get('follow_up_documented')}`")
            st.write(f"**Completeness:** `{audit.get('overall_completeness')}`")
        if audit.get("missing_elements"): st.warning("Missing:"); [st.write(f"• {m}") for m in audit["missing_elements"]]
        if audit.get("priority_actions"): st.error("Priority:"); [st.write(f"• {a}") for a in audit["priority_actions"]]
        if audit.get("standard_recommendations"): [st.write(f"• {r}") for r in audit["standard_recommendations"]]

    st.divider(); st.subheader("⑤ Clinician Sign-off")
    st.warning("All AI outputs require clinician review before filing.")
    cn=st.text_input("Clinician Name"); dg=st.text_input("Designation / MMC No.")
    sd=st.date_input("Review Date", value=datetime.now())
    if st.checkbox("✅ I have reviewed all outputs and approve filing") and cn:
        st.success(f"✅ Approved by {cn} ({dg}) on {sd}"); st.balloons()

st.divider(); st.subheader("⑥ Feedback — Validity & Feasibility")
st.caption("No patient data collected. Sent to Dr Naim.")
with st.form("fb"):
    fn=st.text_input("Name (optional)")
    fr=st.selectbox("Role",["Psychiatrist","Clinical Psychologist","Counsellor","Medical Officer","Nurse","Other"])
    fv=st.slider("Validity (1=not valid, 5=highly valid)",1,5,3)
    ff=st.slider("Feasibility (1=not feasible, 5=highly feasible)",1,5,3)
    fu=st.multiselect("Most useful:",["Audio transcription","Image extraction","Thematic analysis","Clinical note","ICD-11/DSM-5","Recommendations","Doc quality checker","Sign-off"])
    fi=st.text_area("Suggestions",height=80); fc=st.text_area("Comments",height=60)
    if st.form_submit_button("📨 Submit Feedback", use_container_width=True):
        try:
            rp=_req.post("https://api.emailjs.com/api/v1.0/email/send",
                json={"service_id":"service_z0rh03y","template_id":"template_6bzqjul","user_id":"jRs1TWPZym8iUjNck",
                    "template_params":{"to_email":"ppnaim@moh.gov.my","name":fn or "Anonymous","role":fr,
                    "validity":str(fv),"feasibility":str(ff),"useful":", ".join(fu) if fu else "None",
                    "improvements":fi or "None","comments":fc or "None","timestamp":datetime.now().strftime("%Y-%m-%d %H:%M")}},
                headers={"Content-Type":"application/json"},timeout=10)
            st.success("✅ Feedback sent to ppnaim@moh.gov.my") if rp.status_code==200 else st.warning(f"Failed: {rp.text}")
        except Exception as e: st.warning(f"Failed: {e}")

st.divider()
st.markdown("<div style='text-align:center;color:#888;font-size:0.75em;'>🧠 <strong>TherapyAI Demo</strong> · Dr Naim AI Team · HTPN · 2026<br><em>No patient records stored. Clinician verification required.</em></div>", unsafe_allow_html=True)
